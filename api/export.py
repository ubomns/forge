from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import StreamingResponse
from slowapi import Limiter

from api.schema import BusinessPlan, SECTION_DISPLAY_NAMES, SECTION_ORDER


def _add_heading(doc: Document, text: str, level: int = 2) -> None:
    doc.add_heading(text, level=level)


def _add_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    doc.add_paragraph(text, style=style)


def _add_bullet_list(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _build_docx(plan: BusinessPlan) -> bytes:
    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x1C, 0x19, 0x17)

    title = doc.add_heading("Business Plan", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    tagline = plan.executive_summary.tagline
    if tagline:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(tagline)
        run.italic = True
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x78, 0x71, 0x6C)

    doc.add_paragraph()

    for section_id in SECTION_ORDER:
        section_data = getattr(plan, section_id, None)
        if section_data is None:
            continue

        display_name = SECTION_DISPLAY_NAMES.get(section_id, section_id.replace("_", " ").title())
        _add_heading(doc, display_name, level=2)
        data = section_data.model_dump()

        field_labels: dict[str, str] = {
            "tagline": "Tagline",
            "elevator_pitch": "Elevator Pitch",
            "description": None,
            "pain_points": "Pain Points",
            "key_features": "Key Features",
            "tam": "Total Addressable Market",
            "sam": "Serviceable Addressable Market",
            "som": "Serviceable Obtainable Market",
            "segments": "Target Segments",
            "overview": "Landscape",
            "advantages": "Advantages",
            "competitors": "Known Competitors",
            "strategy": "Strategy",
            "channels": "Channels",
            "milestones": "Milestones",
            "revenue_streams": "Revenue Streams",
            "pricing": "Pricing",
            "unit_economics": "Unit Economics",
            "needed_roles": "Key Roles",
            "current_gaps": "Current Gaps",
            "revenue_projection": "Revenue Projection",
            "key_costs": "Key Costs",
            "break_even": "Break-even",
            "amount": "Amount Sought",
            "use_of_funds": "Use of Funds",
            "vision": "Vision",
        }

        for field_name, value in data.items():
            if value is None:
                continue
            label = field_labels.get(field_name)
            if isinstance(value, list):
                if label and label != field_name.replace("_", " ").title():
                    _add_heading(doc, label, level=3)
                _add_bullet_list(doc, value)
            else:
                if label and label != field_name.replace("_", " ").title():
                    _add_heading(doc, label, level=3)
                _add_paragraph(doc, str(value))

        doc.add_paragraph()

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def register_export_routes(app: FastAPI, limiter: Limiter) -> None:
    @app.post("/api/export")
    @limiter.limit("10/minute")
    async def export_docx(request: Request, body: dict[str, Any]):
        try:
            plan = BusinessPlan.model_validate(body)
        except Exception as exc:
            raise HTTPException(
                status_code=422,
                detail={"error": {"code": "INVALID_REQUEST", "message": f"Invalid plan data: {exc}"}},
            )

        try:
            docx_bytes = _build_docx(plan)
        except Exception as exc:
            raise HTTPException(
                status_code=500,
                detail={"error": {"code": "UPSTREAM_ERROR", "message": f"Failed to generate document: {exc}"}},
            )

        return StreamingResponse(
            io.BytesIO(docx_bytes),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": 'attachment; filename="forge-business-plan.docx"'},
        )